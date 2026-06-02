#!/usr/bin/env python3
"""Generate scenarios/test-scenarios.xlsx — the editable test-case template.

Run this once to (re)create the Excel template. It will NOT overwrite an existing
file unless you pass --force, so you never lose your own scenarios by accident.

    python3 scripts/generate_scenarios_template.py [--force]

Columns:
  ID                  unique id, e.g. TC-001
  Scenario            short name of the test case
  Category            grouping, e.g. Greeting / Orders / Refund / Escalation
  Priority            High / Medium / Low
  Preconditions       anything that must be true before the test
  User Messages       the message(s) to send the bot. For a multi-turn dialog,
                      put each turn on its own line inside the cell
                      (Alt+Enter in Excel) OR separate turns with " || ".
  Expected Result     intended bot behaviour (described, not a verbatim string)
  Notes               free text
"""
import argparse
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OUT = os.path.join(ROOT, "scenarios", "test-scenarios.xlsx")

HEADERS = [
    ("ID", 12),
    ("Scenario", 32),
    ("Category", 16),
    ("Priority", 10),
    ("Preconditions", 28),
    ("User Messages", 46),
    ("Expected Result", 48),
    ("Notes", 26),
]

# Example scenarios for a generic Jira Service Desk support bot.
# Replace / extend these with your real test cases.
EXAMPLES = [
    ["TC-001", "Greeting / smoke test", "Greeting", "High", "Test conversation is open",
     "Hi",
     "Bot replies with a greeting and offers help / shows it is online.", "Basic smoke check"],
    ["TC-002", "Ask what the bot can do", "Capabilities", "Medium", "",
     "What can you help me with?",
     "Bot explains its scope / lists the topics or actions it supports.", ""],
    ["TC-003", "Order status request", "Orders", "High", "",
     "Where is my order?",
     "Bot asks for an order id / details, or explains how to check order status.", ""],
    ["TC-004", "Refund request", "Refund", "High", "",
     "I want a refund for my last order",
     "Bot starts a refund flow or explains the refund policy / next steps.", ""],
    ["TC-005", "Multi-turn: order then refund", "Orders", "Medium", "",
     "I have a problem with my order || It arrived damaged || Can I get a refund?",
     "Bot keeps context across turns and guides toward a resolution / refund.",
     "Each line is one user message"],
    ["TC-006", "Escalate to a human agent", "Escalation", "High", "",
     "I want to talk to a human agent",
     "Bot offers handover to a human / creates a ticket / explains agent availability.", ""],
    ["TC-007", "Out-of-scope question", "Robustness", "Low", "",
     "What's the weather tomorrow?",
     "Bot politely declines / redirects to support topics instead of hallucinating.", ""],
    ["TC-008", "Non-English message", "Localization", "Medium", "",
     "Привет, мне нужна помощь с заказом",
     "Bot understands and responds appropriately (ideally in the same language).", ""],
    ["TC-009", "Empty / nonsense input", "Robustness", "Low", "",
     "asdfghjkl",
     "Bot handles gracefully — asks to rephrase, does not crash or error.", ""],
    ["TC-010", "Goodbye / close conversation", "Greeting", "Low", "",
     "Thanks, bye",
     "Bot acknowledges and closes politely.", ""],
]


def build(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Scenarios"

    header_fill = PatternFill("solid", fgColor="1F6FEB")
    header_font = Font(bold=True, color="FFFFFF")
    wrap_top = Alignment(wrap_text=True, vertical="top")

    for col, (name, width) in enumerate(HEADERS, start=1):
        c = ws.cell(row=1, column=col, value=name)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width

    for r, row in enumerate(EXAMPLES, start=2):
        for col, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.alignment = wrap_top

    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22

    # A small "How to use" note sheet
    notes = wb.create_sheet("README")
    tips = [
        "How to write test scenarios:",
        "",
        "- One row = one test scenario.",
        "- ID must be unique (TC-001, TC-002, ...).",
        "- For a multi-turn dialog, put each user message on its own line in the",
        "  'User Messages' cell (Alt+Enter), or separate turns with ' || '.",
        "- 'Expected Result' describes intended behaviour in plain language — the",
        "  test agent judges by meaning, not exact text match.",
        "- Leave 'Notes' for anything a human reviewer should know.",
        "",
        "The example rows on the 'Scenarios' sheet are safe to delete or replace.",
    ]
    for i, line in enumerate(tips, start=1):
        notes.cell(row=i, column=1, value=line)
    notes.column_dimensions["A"].width = 80

    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--force", action="store_true",
                    help="overwrite an existing scenarios file")
    args = ap.parse_args()

    if os.path.exists(OUT) and not args.force:
        print(f"[skip] {OUT} already exists. Use --force to overwrite.")
        return
    build(OUT)
    print(f"[ok] wrote {OUT}")


if __name__ == "__main__":
    sys.exit(main())
