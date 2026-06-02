#!/usr/bin/env python3
"""Turn reports/results.json (written by the agent) into an Excel report.

    python3 scripts/report.py [reports/results.json]

Writes reports/report-<timestamp>.xlsx and prints a PASS/FAIL/BLOCKED summary.
"""
import datetime as dt
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DEFAULT_RESULTS = os.path.join(ROOT, "reports", "results.json")

STATUS_COLORS = {
    "PASS": "1A7F37",
    "FAIL": "CF222E",
    "BLOCKED": "9A6700",
}

COLUMNS = [
    ("ID", 12),
    ("Scenario", 30),
    ("Category", 16),
    ("Priority", 10),
    ("Status", 12),
    ("Expected", 44),
    ("Actual", 44),
    ("Dialog (turns)", 50),
    ("Screenshot", 28),
    ("Notes", 26),
]


def fmt_turns(turns):
    out = []
    for t in turns or []:
        u = t.get("user", "")
        b = t.get("bot", "")
        out.append(f"USER: {u}\nBOT: {b}")
    return "\n---\n".join(out)


def main():
    results_path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_RESULTS
    if not os.path.exists(results_path):
        print(f"[error] results file not found: {results_path}", file=sys.stderr)
        print("The test agent should write reports/results.json first.", file=sys.stderr)
        return 1

    with open(results_path, encoding="utf-8") as f:
        data = json.load(f)

    results = data.get("results", [])

    wb = Workbook()
    ws = wb.active
    ws.title = "Results"

    header_fill = PatternFill("solid", fgColor="1F6FEB")
    header_font = Font(bold=True, color="FFFFFF")
    wrap_top = Alignment(wrap_text=True, vertical="top")

    # Meta block
    ws.cell(row=1, column=1, value="Support Bot Test Report").font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value=f"Target: {data.get('target_url', '')}")
    ws.cell(row=3, column=1, value=f"Started: {data.get('run_started_at', '')}   "
                                   f"Finished: {data.get('run_finished_at', '')}")

    counts = {"PASS": 0, "FAIL": 0, "BLOCKED": 0}
    for r in results:
        st = str(r.get("status", "")).upper()
        if st in counts:
            counts[st] += 1
    ws.cell(row=4, column=1,
            value=f"Total: {len(results)}   PASS: {counts['PASS']}   "
                  f"FAIL: {counts['FAIL']}   BLOCKED: {counts['BLOCKED']}").font = Font(bold=True)

    header_row = 6
    for col, (name, width) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=header_row, column=col, value=name)
        c.fill = header_fill
        c.font = header_font
        ws.column_dimensions[get_column_letter(col)].width = width

    for i, r in enumerate(results):
        row = header_row + 1 + i
        st = str(r.get("status", "")).upper()
        values = [
            r.get("id", ""),
            r.get("name", ""),
            r.get("category", ""),
            r.get("priority", ""),
            st,
            r.get("expected", ""),
            r.get("actual", ""),
            fmt_turns(r.get("turns")),
            r.get("screenshot", ""),
            r.get("notes", ""),
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.alignment = wrap_top
        status_cell = ws.cell(row=row, column=5)
        if st in STATUS_COLORS:
            status_cell.font = Font(bold=True, color=STATUS_COLORS[st])

    ws.freeze_panes = f"A{header_row + 1}"

    ts = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    out = os.path.join(ROOT, "reports", f"report-{ts}.xlsx")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)

    print(f"[ok] wrote {out}")
    print(f"Total: {len(results)}  PASS: {counts['PASS']}  "
          f"FAIL: {counts['FAIL']}  BLOCKED: {counts['BLOCKED']}")
    for r in results:
        st = str(r.get("status", "")).upper()
        if st in ("FAIL", "BLOCKED"):
            print(f"  [{st}] {r.get('id', '')} {r.get('name', '')}: "
                  f"{r.get('notes') or r.get('actual', '')}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
