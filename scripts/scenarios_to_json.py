#!/usr/bin/env python3
"""Read the Excel test cases and emit scenarios/scenarios.json for the agent.

    python3 scripts/scenarios_to_json.py [path/to/file.xlsx]

Multi-turn dialogs: a 'User Messages' cell may contain several turns, separated
either by newlines or by the ' || ' delimiter. Each becomes one entry in `turns`.
"""
import json
import os
import sys

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)

DEFAULT_XLSX = os.path.join(ROOT, "scenarios", "test-scenarios.xlsx")
OUT_JSON = os.path.join(ROOT, "scenarios", "scenarios.json")

# Map normalized header -> canonical field name
HEADER_MAP = {
    "id": "id",
    "scenario": "name",
    "category": "category",
    "priority": "priority",
    "preconditions": "preconditions",
    "user messages": "user_messages",
    "expected result": "expected_result",
    "notes": "notes",
}


def split_turns(raw):
    if raw is None:
        return []
    text = str(raw)
    # Prefer explicit '||' delimiter, else split on newlines.
    if "||" in text:
        parts = text.split("||")
    else:
        parts = text.splitlines()
    return [p.strip() for p in parts if p and p.strip()]


def main():
    xlsx = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    if not os.path.exists(xlsx):
        print(f"[error] scenarios file not found: {xlsx}", file=sys.stderr)
        print("Run: python3 scripts/generate_scenarios_template.py", file=sys.stderr)
        return 1

    wb = load_workbook(xlsx, data_only=True)
    ws = wb["Scenarios"] if "Scenarios" in wb.sheetnames else wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        print(f"[error] empty sheet in {xlsx}", file=sys.stderr)
        return 1

    header = rows[0]
    # Build column index -> field
    fields = {}
    for idx, name in enumerate(header):
        if name is None:
            continue
        key = str(name).strip().lower()
        if key in HEADER_MAP:
            fields[idx] = HEADER_MAP[key]

    if "id" not in fields.values():
        print("[error] could not find an 'ID' column in the header row.", file=sys.stderr)
        return 1

    scenarios = []
    for row in rows[1:]:
        if row is None:
            continue
        record = {}
        for idx, field in fields.items():
            record[field] = row[idx] if idx < len(row) else None
        # Skip blank rows (no id and no messages)
        if not record.get("id") and not record.get("user_messages"):
            continue
        turns = split_turns(record.get("user_messages"))
        scenarios.append({
            "id": (str(record.get("id")).strip() if record.get("id") else ""),
            "name": (str(record.get("name")).strip() if record.get("name") else ""),
            "category": (str(record.get("category")).strip() if record.get("category") else ""),
            "priority": (str(record.get("priority")).strip() if record.get("priority") else ""),
            "preconditions": (str(record.get("preconditions")).strip() if record.get("preconditions") else ""),
            "turns": turns,
            "expected_result": (str(record.get("expected_result")).strip() if record.get("expected_result") else ""),
            "notes": (str(record.get("notes")).strip() if record.get("notes") else ""),
        })

    os.makedirs(os.path.dirname(OUT_JSON), exist_ok=True)
    with open(OUT_JSON, "w", encoding="utf-8") as f:
        json.dump(scenarios, f, ensure_ascii=False, indent=2)

    print(f"[ok] {len(scenarios)} scenario(s) -> {OUT_JSON}")
    for s in scenarios:
        print(f"  - {s['id']}: {s['name']} ({len(s['turns'])} turn(s))")
    return 0


if __name__ == "__main__":
    sys.exit(main())
